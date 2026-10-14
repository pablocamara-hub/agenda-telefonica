"""Agenda telefónica v1 (rama con-excel).

Programa de consola para gestionar contactos.
Cada contacto tiene nombre y teléfono.
Los datos se guardan en una hoja de cálculo: data/contactos.xlsx.
"""

import os

from openpyxl import Workbook, load_workbook

RUTA_DATOS = "data/contactos.xlsx"
NOMBRE_HOJA = "Contactos"
CABECERA_NOMBRE = "Nombre"
CABECERA_TELEFONO = "Telefono"


def telefono_es_valido(telefono):
    """Comprueba que el teléfono no esté vacío y solo use dígitos, espacios, + y -."""
    texto = telefono.strip()
    if texto == "":
        return False
    permitidos = set("0123456789 +-")
    for letra in texto:
        if letra not in permitidos:
            return False
    return True


def anadir_contacto(contactos, nombre, telefono):
    """Añade un contacto a la lista.

    Devuelve (True, mensaje) si se añade y (False, mensaje) si falla.
    No acepta nombres vacíos, teléfonos no válidos ni nombres duplicados.
    """
    nombre_limpio = nombre.strip()
    telefono_limpio = telefono.strip()
    if nombre_limpio == "":
        return False, "El nombre no puede estar vacío."
    if not telefono_es_valido(telefono_limpio):
        return False, "El teléfono no es válido."
    for contacto in contactos:
        if contacto["nombre"] == nombre_limpio:
            return False, "Ya existe un contacto con ese nombre."
    contactos.append({"nombre": nombre_limpio, "telefono": telefono_limpio})
    return True, "Contacto añadido."


def buscar_por_nombre(contactos, texto):
    """Busca contactos por nombre parcial, sin distinguir mayúsculas."""
    buscado = texto.strip().lower()
    encontrados = []
    for contacto in contactos:
        if buscado in contacto["nombre"].lower():
            encontrados.append(contacto)
    return encontrados


def buscar_por_numero(contactos, texto):
    """Busca contactos por número parcial o exacto."""
    buscado = texto.strip()
    encontrados = []
    for contacto in contactos:
        if buscado in contacto["telefono"]:
            encontrados.append(contacto)
    return encontrados
def eliminar_contacto(contactos, nombre):
    """Elimina el contacto con el nombre exacto. La confirmación se pide en el menú."""
    buscado = nombre.strip()
    for i in range(len(contactos)):
        if contactos[i]["nombre"] == buscado:
            del contactos[i]
            return True
    return False


def ordenar_contactos(contactos, inverso=False):
    """Devuelve una lista nueva ordenada por nombre sin distinguir mayúsculas."""
    return sorted(contactos, key=lambda c: c["nombre"].lower(), reverse=inverso)


def cargar_contactos(ruta=RUTA_DATOS):
    """Carga los contactos desde la hoja de cálculo. Si no existe, devuelve lista vacía."""
    contactos = []
    if not os.path.exists(ruta):
        return contactos
    libro = load_workbook(ruta)
    if NOMBRE_HOJA not in libro.sheetnames:
        print("Aviso: hoja no encontrada: " + NOMBRE_HOJA)
        return contactos
    hoja = libro[NOMBRE_HOJA]
    primera = True
    for fila in hoja.iter_rows(values_only=True):
        if primera:
            primera = False
            continue
        if fila is None:
            continue
        nombre = fila[0] if len(fila) > 0 else None
        telefono = fila[1] if len(fila) > 1 else None
        if nombre is None and telefono is None:
            continue
        nombre = str(nombre).strip() if nombre is not None else ""
        telefono = str(telefono).strip() if telefono is not None else ""
        if nombre == "" or not telefono_es_valido(telefono):
            print("Aviso: fila ignorada: " + nombre + ";" + telefono)
            continue
        contactos.append({"nombre": nombre, "telefono": telefono})
    return contactos


def guardar_contactos(contactos, ruta=RUTA_DATOS):
    """Guarda los contactos en la hoja de cálculo con cabecera Nombre/Telefono."""
    carpeta = os.path.dirname(ruta)
    if carpeta != "" and not os.path.exists(carpeta):
        os.makedirs(carpeta)
    libro = Workbook()
    hoja = libro.active
    hoja.title = NOMBRE_HOJA
    hoja["A1"] = CABECERA_NOMBRE
    hoja["B1"] = CABECERA_TELEFONO
    fila_numero = 2
    for contacto in contactos:
        hoja.cell(row=fila_numero, column=1, value=contacto["nombre"]).number_format = "@"
        hoja.cell(row=fila_numero, column=2, value=contacto["telefono"]).number_format = "@"
        fila_numero = fila_numero + 1
    libro.save(ruta)


def mostrar_contactos(contactos, inverso=False):
    """Muestra los contactos ordenados por pantalla."""
    ordenados = ordenar_contactos(contactos, inverso)
    if len(ordenados) == 0:
        print("No hay contactos.")
        return
    for contacto in ordenados:
        print(contacto["nombre"] + ": " + contacto["telefono"])


def mostrar_menu():
    """Muestra las opciones del menú."""
    print("\n--- AGENDA ---")
    print("1. Añadir contacto")
    print("2. Buscar por nombre")
    print("3. Buscar por número")
    print("4. Eliminar contacto")
    print("5. Mostrar A-Z")
    print("6. Mostrar Z-A")
    print("0. Guardar y salir")


def main():
    """Bucle principal del programa."""
    contactos = cargar_contactos()
    while True:
        mostrar_menu()
        opcion = input("Elige una opción: ").strip()
        if opcion == "1":
            nombre = input("Nombre: ")
            telefono = input("Teléfono: ")
            _, mensaje = anadir_contacto(contactos, nombre, telefono)
            print(mensaje)
        elif opcion == "2":
            texto = input("Nombre a buscar: ")
            encontrados = buscar_por_nombre(contactos, texto)
            mostrar_contactos(encontrados)
        elif opcion == "3":
            texto = input("Número a buscar: ")
            encontrados = buscar_por_numero(contactos, texto)
            mostrar_contactos(encontrados)
        elif opcion == "4":
            nombre = input("Nombre exacto a eliminar: ")
            respuesta = input("¿Seguro? (s/n): ").strip().lower()
            if respuesta == "s":
                if eliminar_contacto(contactos, nombre):
                    print("Contacto eliminado.")
                else:
                    print("No existe ese contacto.")
            else:
                print("Cancelado.")
        elif opcion == "5":
            mostrar_contactos(contactos, inverso=False)
        elif opcion == "6":
            mostrar_contactos(contactos, inverso=True)
        elif opcion == "0":
            guardar_contactos(contactos)
            print("Agenda guardada. ¡Adiós!")
            break
        else:
            print("Opción no válida.")


if __name__ == "__main__":
    main()
