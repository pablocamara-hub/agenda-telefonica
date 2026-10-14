"""Pruebas de la agenda telefónica v1 (rama con-excel)."""

import os
import tempfile

from openpyxl import Workbook, load_workbook

from src.agenda import (
    CABECERA_NOMBRE,
    CABECERA_TELEFONO,
    NOMBRE_HOJA,
    anadir_contacto,
    buscar_por_nombre,
    buscar_por_numero,
    cargar_contactos,
    eliminar_contacto,
    guardar_contactos,
    ordenar_contactos,
    telefono_es_valido,
)


def test_telefono_valido():
    assert telefono_es_valido("612345678") is True
    assert telefono_es_valido("+34 612-345-678") is True
    assert telefono_es_valido("") is False
    assert telefono_es_valido("abc") is False


def test_anadir_ok():
    contactos = []
    ok, _ = anadir_contacto(contactos, "Ana", "612345678")
    assert ok is True
    assert len(contactos) == 1


def test_anadir_rechaza_vacio_y_duplicado():
    contactos = []
    ok, _ = anadir_contacto(contactos, "", "123")
    assert ok is False
    ok, _ = anadir_contacto(contactos, "Ana", "")
    assert ok is False
    anadir_contacto(contactos, "Ana", "111")
    ok, _ = anadir_contacto(contactos, "Ana", "222")
    assert ok is False
    assert len(contactos) == 1


def test_buscar_por_nombre_parcial():
    contactos = [{"nombre": "Ana", "telefono": "111"}]
    assert len(buscar_por_nombre(contactos, "ana")) == 1
    assert len(buscar_por_nombre(contactos, "AN")) == 1
    assert len(buscar_por_nombre(contactos, "Luis")) == 0


def test_buscar_por_numero():
    contactos = [{"nombre": "Ana", "telefono": "612345678"}]
    assert len(buscar_por_numero(contactos, "612")) == 1
    assert len(buscar_por_numero(contactos, "999")) == 0


def test_eliminar():
    contactos = [{"nombre": "Ana", "telefono": "111"}]
    assert eliminar_contacto(contactos, "Ana") is True
    assert contactos == []
    assert eliminar_contacto(contactos, "Ana") is False


def test_ordenar():
    contactos = [
        {"nombre": "Luis", "telefono": "2"},
        {"nombre": "ana", "telefono": "1"},
    ]
    az = [c["nombre"] for c in ordenar_contactos(contactos)]
    assert az == ["ana", "Luis"]
    za = [c["nombre"] for c in ordenar_contactos(contactos, inverso=True)]
    assert za == ["Luis", "ana"]


def test_guardar_y_cargar():
    contactos = [{"nombre": "Ana", "telefono": "111"}]
    with tempfile.TemporaryDirectory() as tmp:
        ruta = os.path.join(tmp, "contactos.xlsx")
        guardar_contactos(contactos, ruta)
        cargados = cargar_contactos(ruta)
        assert cargados == contactos


def test_guardar_crea_cabecera():
    with tempfile.TemporaryDirectory() as tmp:
        ruta = os.path.join(tmp, "contactos.xlsx")
        guardar_contactos([{"nombre": "Ana", "telefono": "111"}], ruta)
        libro = load_workbook(ruta)
        assert NOMBRE_HOJA in libro.sheetnames
        hoja = libro[NOMBRE_HOJA]
        assert hoja["A1"].value == CABECERA_NOMBRE
        assert hoja["B1"].value == CABECERA_TELEFONO
        assert hoja["A2"].value == "Ana"
        assert str(hoja["B2"].value) == "111"


def test_cargar_ignora_filas_malas():
    with tempfile.TemporaryDirectory() as tmp:
        ruta = os.path.join(tmp, "contactos.xlsx")
        libro = Workbook()
        hoja = libro.active
        hoja.title = NOMBRE_HOJA
        hoja["A1"] = CABECERA_NOMBRE
        hoja["B1"] = CABECERA_TELEFONO
        hoja.append(["Ana", "111"])
        hoja.append(["", "222"])
        hoja.append(["Luis", "abc"])
        libro.save(ruta)
        cargados = cargar_contactos(ruta)
        assert cargados == [{"nombre": "Ana", "telefono": "111"}]
